# -*- coding: utf-8 -*-
"""签到结果导出：生成 .xlsx，含「签到明细」与「签到列表」两张表。"""
import os
from dataclasses import dataclass
from datetime import datetime

from .config import DATA_DIR


@dataclass
class CheckinRecord:
    """一条签到记录。"""
    name: str                # 姓名
    class_name: str          # 班级
    time_str: str            # 签到时间（HH:MM:SS）
    result: str              # 结果：成功 / 失败 / 重复
    similarity: float = 0.0  # 相似度
    student_id: str = ""     # 身份证号
    school: str = ""         # 学校
    grade: str = ""          # 年级
    level: str = ""          # 匹配级别：class / grade


_HEADERS = ["序号", "姓名", "班级", "签到时间", "结果", "相似度", "身份证号", "学校", "年级", "匹配级别"]
_RESULT_COLORS = {"成功": "16A34A", "失败": "DC2626", "重复": "F59E0B"}

# 签到列表（花名册）：范围内每个学生 + 状态 + 签到时间
_ROSTER_HEADERS = ["序号", "学校", "年级", "班级", "姓名", "身份证号", "状态", "签到时间"]
_ROSTER_COLORS = {"已签到": "16A34A", "未签到": "DC2626"}


def _load_openpyxl():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        return Workbook, Alignment, Font, PatternFill
    except ImportError as e:
        raise RuntimeError("缺少 openpyxl 依赖，请先双击「安装环境.bat」安装。") from e


def _header_style(cell, PatternFill, Font, Alignment):
    cell.fill = PatternFill("solid", fgColor="2563EB")
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _write_detail_sheet(ws, records, Workbook, Alignment, Font, PatternFill):
    """签到明细：逐条签到记录。"""
    ws.title = "签到明细"
    for col, h in enumerate(_HEADERS, 1):
        _header_style(ws.cell(row=1, column=col, value=h), PatternFill, Font, Alignment)

    for i, r in enumerate(records, 2):
        row = [
            i - 1, r.name, r.class_name, r.time_str, r.result,
            round(r.similarity, 3), r.student_id, r.school, r.grade, r.level,
        ]
        for col, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=col, value=v)
            if r.result in _RESULT_COLORS and col == 5:
                cell.font = Font(color=_RESULT_COLORS[r.result], bold=True)

    widths = [6, 12, 16, 12, 10, 10, 24, 16, 12, 10]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = w
    ws.freeze_panes = "A2"


def _write_roster_sheet(ws, roster, Workbook, Alignment, Font, PatternFill):
    """签到列表：范围内每个学生 + 是否已签到 + 签到时间。"""
    ws.title = "签到列表"
    for col, h in enumerate(_ROSTER_HEADERS, 1):
        _header_style(ws.cell(row=1, column=col, value=h), PatternFill, Font, Alignment)

    for i, s in enumerate(roster, 2):
        row = [
            i - 1, s["school"], s["grade"], s["cls"], s["name"],
            s["student_id"], s["status"], s["sign_time"],
        ]
        for col, v in enumerate(row, 1):
            cell = ws.cell(row=i, column=col, value=v)
            if s["status"] in _ROSTER_COLORS and col == 7:
                cell.font = Font(color=_ROSTER_COLORS[s["status"]], bold=True)

    widths = [6, 16, 12, 16, 12, 24, 10, 12]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = w
    ws.freeze_panes = "A2"


def make_export_path(scope_name: str = "") -> str:
    """生成带时间戳的导出文件路径（保存到 data 目录，gitignore 排除）。"""
    os.makedirs(DATA_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = f"签到结果_{scope_name}_{ts}.xlsx" if scope_name else f"签到结果_{ts}.xlsx"
    return os.path.join(DATA_DIR, name)


def export_checkin(records, roster, filepath: str = None) -> str:
    """导出两张表：签到明细（逐条）+ 签到列表（范围内每个学生的签到状态）。
    roster: [{school, grade, cls, name, student_id, status, sign_time}, ...]
    返回文件路径。"""
    Workbook, Alignment, Font, PatternFill = _load_openpyxl()

    if filepath is None:
        filepath = make_export_path()

    wb = Workbook()
    ws_detail = wb.active
    _write_detail_sheet(ws_detail, records, Workbook, Alignment, Font, PatternFill)
    ws_roster = wb.create_sheet()
    _write_roster_sheet(ws_roster, roster, Workbook, Alignment, Font, PatternFill)

    wb.save(filepath)
    return filepath


def export_xlsx(records, filepath: str = None) -> str:
    """仅导出「签到明细」单表（兼容旧调用，如 test_export.py）。"""
    Workbook, Alignment, Font, PatternFill = _load_openpyxl()

    if filepath is None:
        filepath = make_export_path()

    wb = Workbook()
    ws = wb.active
    _write_detail_sheet(ws, records, Workbook, Alignment, Font, PatternFill)
    wb.save(filepath)
    return filepath
