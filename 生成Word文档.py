# -*- coding: utf-8 -*-
"""把「使用说明书.md」「免责声明.md」转成 Word 文档（.docx）。"""
import re

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

FONT = "微软雅黑"
BLUE = (0x25, 0x63, 0xEB)
DARK = (0x0F, 0x17, 0x2A)
GRAY = (0x47, 0x55, 0x69)


def _style_run(run, text, bold=False, size=12, color=None):
    run.text = text
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def _add_heading(doc, text, level):
    p = doc.add_paragraph()
    if level == 1:
        size, color = 22, BLUE
    elif level == 2:
        size, color = 16, BLUE
    else:
        size, color = 13, DARK
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(8)
    _style_run(p.add_run(), text, bold=True, size=size, color=color)


def _add_para(doc, text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.4
    # 处理 **加粗** 和 `代码`
    for tok in re.split(r"(\*\*.*?\*\*|`[^`]*`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            _style_run(p.add_run(), tok[2:-2], bold=True, size=size)
        elif tok.startswith("`") and tok.endswith("`"):
            _style_run(p.add_run(), tok[1:-1], size=size, color=(0xC2, 0x41, 0x0C))
        else:
            _style_run(p.add_run(), tok, size=size)


def _add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    for tok in re.split(r"(\*\*.*?\*\*|`[^`]*`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            _style_run(p.add_run(), tok[2:-2], bold=True)
        elif tok.startswith("`") and tok.endswith("`"):
            _style_run(p.add_run(), tok[1:-1], color=(0xC2, 0x41, 0x0C))
        else:
            _style_run(p.add_run(), tok)


def _add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = table.cell(i, j)
            c.text = ""
            _style_run(c.paragraphs[0].add_run(), cell.strip(), bold=(i == 0), size=10.5)


def _split_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("### "):
            _add_heading(doc, line[4:], 3)
        elif line.startswith("## "):
            _add_heading(doc, line[3:], 2)
        elif line.startswith("# "):
            _add_heading(doc, line[2:], 1)
        elif line.startswith("|"):
            # 收集连续表格行
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = _split_table_row(lines[i])
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    i += 1  # 跳过分隔行
                    continue
                rows.append(cells)
                i += 1
            if rows:
                _add_table(doc, rows)
            continue
        elif line.startswith("- ") or line.startswith("* "):
            _add_bullet(doc, line[2:])
        elif line == "---":
            pass
        elif line.startswith("> "):
            _add_para(doc, line[2:], size=11)
        else:
            _add_para(doc, line)
        i += 1

    doc.save(docx_path)
    print(f"已生成：{docx_path}")


if __name__ == "__main__":
    convert("使用说明书.md", "使用说明书.docx")
    convert("免责声明.md", "免责声明.docx")
    print("完成。")
